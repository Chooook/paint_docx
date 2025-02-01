import gc
from functools import lru_cache
from typing import Iterator, Tuple

from docx import Document
from docx.text.run import Run


class RunCollector:
    def __init__(self, document: Document) -> None:
        @lru_cache(maxsize=1)
        def create_space_runs() -> Tuple[Run, Run]:
            # Для создания Run вне текущего документа, нужен объект Document
            temp_doc = Document()
            return (temp_doc.add_paragraph().add_run(' '),
                    temp_doc.add_paragraph().add_run('\n'))

        self.__space_run, self.__line_brake_run = create_space_runs()
        self.document = document
        self.document_runs = None

        self.text_runs = self.update_text_runs()
        self.tables_runs_by_rows = self.__update_tables_runs_by_rows()
        self.tables_runs_by_columns = self.__update_tables_runs_by_columns()
        self.footnotes_runs = self.update_footnotes_runs()

        self.document_runs = self.update_document_runs(reload_all=False)

    def __iter__(self) -> Iterator[Run]:
        return iter(self.document_runs)

    def update_text_runs(self) -> list[Run]:
        runs = []
        for paragraph in self.document.paragraphs:
            runs.extend(paragraph.runs)
            runs.extend([self.__space_run])
        self.text_runs = runs
        if self.document_runs:
            self.update_document_runs(reload_all=False)
        return self.text_runs

    def update_tables_runs(self) -> tuple[list[Run], list[Run]]:
        # В случае объединённых ячеек данные дублируются,
        # т.к. фигурируют "в нескольких колонках или строках".
        # Не баг, а фича, т.к. позволяет отследить различные варианты
        # поколоночных и построчных сочетаний с объединёнными ячейками.
        self.tables_runs_by_rows = self.__update_tables_runs_by_rows()
        self.tables_runs_by_columns = self.__update_tables_runs_by_columns()
        if self.document_runs:
            self.update_document_runs(reload_all=False)
        return self.tables_runs_by_rows, self.tables_runs_by_columns

    def __update_tables_runs_by_rows(self) -> list[Run]:
        runs = []
        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        runs.extend(paragraph.runs)
                        runs.extend([self.__space_run])
        return runs

    def __update_tables_runs_by_columns(self) -> list[Run]:
        runs = []
        for table in self.document.tables:
            num_rows = len(table.rows)
            num_cols = len(table.columns)
            for col in range(num_cols):
                for row in range(num_rows):
                    cell = table.cell(row, col)
                    for paragraph in cell.paragraphs:
                        runs.extend(paragraph.runs)
                        runs.extend([self.__space_run])
        return runs

    def update_footnotes_runs(self) -> list[Run]:
        try:
            # Атрибут footnotes в Document существует не всегда
            document_footnotes = self.document.footnotes
        except AttributeError:
            document_footnotes = []
        runs = []
        for footnote in document_footnotes:
            for paragraph in footnote.paragraphs:
                runs.extend(paragraph.runs)
                runs.extend([self.__space_run])
        self.footnotes_runs = runs
        if self.document_runs:
            self.update_document_runs(reload_all=False)
        return self.footnotes_runs

    def update_document_runs(self, reload_all: bool = True) -> list[Run]:
        if reload_all:
            self.update_text_runs()
            self.update_tables_runs()
            self.update_footnotes_runs()
        # Для разделения объектов Run в различных сущностях документа
        del self.document_runs
        self.document_runs = (self.text_runs
                              + [self.__line_brake_run]
                              + self.tables_runs_by_rows
                              + [self.__line_brake_run]
                              + self.tables_runs_by_columns
                              + [self.__line_brake_run]
                              + self.footnotes_runs)
        # Очистка памяти необходима, т.к. при работе с изменяемыми объектами,
        # в памяти остаются неиспользуемые ссылки, можно проверить с помощью:
        # >>> len(gc.get_objects())
        gc.collect()
        return self.document_runs
