__all__ = ['RunsMap', 'RunWithSpan']

import gc
from dataclasses import dataclass

from docx import Document
from docx.text.run import Run


@dataclass(frozen=True)
class RunWithSpan:
    start:  int
    end: int
    run: Run


class RunsMap:

    def __init__(self, document: Document):
        self.document = document
        self.__last_span_pos = 0
        self.__document_strings = []

        self.text: str = ''
        self.runs_map: list[RunWithSpan] = []

        self.reload()

    def find_runs_by_span(self, start, end) -> list[RunWithSpan]:
        runs: set[RunWithSpan] = set()
        for run in self.runs_map:
            if (
                run.start <= start < run.end
                or run.start < end <= run.end
                or start < run.start < end
                or start < run.end < end
            ):
                runs.add(run)
        runs: list[RunWithSpan] = list(sorted(runs, key=lambda r: r.start))
        return runs

    def reload(self):
        self.text = ''
        self.runs_map.clear()
        self.__last_span_pos = 0

        # text runs:
        for paragraph in self.document.paragraphs:
            self.__get_runs_and_text(paragraph)
        self.__document_strings += '\n'
        self.__last_span_pos += 1

        # table runs by rows:
        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self.__get_runs_and_text(paragraph)
        self.__document_strings += '\n'
        self.__last_span_pos += 1

        # table runs by columns:
        for table in self.document.tables:
            for col in table.columns:
                for cell in col.cells:
                    for paragraph in cell.paragraphs:
                        self.__get_runs_and_text(paragraph)
        self.__document_strings += '\n'
        self.__last_span_pos += 1

        # footnotes runs:
        try:
            # Атрибут footnotes в Document существует не всегда
            document_footnotes = self.document.footnotes
            for footnote in document_footnotes:
                for paragraph in footnote.paragraphs:
                    self.__get_runs_and_text(paragraph)
        except AttributeError:
            pass

        self.text = ''.join(self.__document_strings)

        self.__document_strings.clear()
        self.__last_span_pos = 0
        gc.collect()


    def __get_runs_and_text(self, paragraph):
        self.__document_strings.append(' ')
        self.__last_span_pos += 1
        paragraph_text = ''.join([run.text for run in paragraph.runs])
        self.__document_strings.append(paragraph_text)

        for run in paragraph.runs:
            run_len = len(run.text)
            self.runs_map.append(
                RunWithSpan(start=self.__last_span_pos,
                            end=self.__last_span_pos + run_len,
                            run=run))
            self.__last_span_pos += run_len
